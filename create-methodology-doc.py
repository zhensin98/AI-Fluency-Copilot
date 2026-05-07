from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page margins ──
section = doc.sections[0]
section.left_margin   = Inches(1.2)
section.right_margin  = Inches(1.2)
section.top_margin    = Inches(1.0)
section.bottom_margin = Inches(1.0)

# ── Helper: set paragraph font ──
def set_font(paragraph, name='Calibri', size=11, bold=False, color=None):
    for run in paragraph.runs:
        run.font.name  = name
        run.font.size  = Pt(size)
        run.font.bold  = bold
        if color:
            run.font.color.rgb = RGBColor(*color)

def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.font.name = 'Calibri'
        run.font.color.rgb = RGBColor(0, 0, 0)
        if level == 1:
            run.font.size = Pt(16)
            run.font.bold = True
        elif level == 2:
            run.font.size = Pt(13)
            run.font.bold = True
        elif level == 3:
            run.font.size = Pt(11)
            run.font.bold = True
    return p

def add_para(doc, text, size=11, bold=False, color=None, space_before=0, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name  = 'Calibri'
    run.font.size  = Pt(size)
    run.font.bold  = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    return p

def add_bullet(doc, text, level=0, size=11):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_divider(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'CCCCCC')
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)

# ═══════════════════════════════════════════════
# TITLE
# ═══════════════════════════════════════════════
p = doc.add_paragraph()
run = p.add_run('Role-Habit Relevance Matrix')
run.font.name  = 'Calibri'
run.font.size  = Pt(22)
run.font.bold  = True
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.paragraph_format.space_after = Pt(4)

p2 = doc.add_paragraph()
run2 = p2.add_run('Methodology & Citation Reference')
run2.font.name  = 'Calibri'
run2.font.size  = Pt(13)
run2.font.bold  = False
run2.font.color.rgb = RGBColor(100, 100, 100)
p2.paragraph_format.space_after = Pt(2)

p3 = doc.add_paragraph()
run3 = p3.add_run('Prepared for: Chanel Microsoft 365 Copilot Value Assessment')
run3.font.name  = 'Calibri'
run3.font.size  = Pt(10)
run3.font.color.rgb = RGBColor(130, 130, 130)
p3.paragraph_format.space_after = Pt(16)

add_divider(doc)

# ═══════════════════════════════════════════════
# SECTION 1 — WHAT IS THE MATRIX
# ═══════════════════════════════════════════════
add_heading(doc, '1.  What Is the Role-Habit Relevance Matrix?', level=1)

add_para(doc,
    'The Role-Habit Relevance Matrix is a prioritisation tool that ranks the seven Microsoft 365 Copilot '
    'usage habits against each organisational role. Rankings run from 1 (most relevant) to 7 (least relevant), '
    'indicating which habits are expected to deliver the greatest productivity benefit for each role.',
    space_after=8)

add_para(doc,
    'Two separate matrices are produced — one for Copilot Chat habits (CH1–CH7) and one for M365 app habits '
    '(MH1–MH7) — covering six organisational roles drawn from the Chanel Profile Explorer:',
    space_after=6)

roles = [
    ('KW',  'Knowledge Workers & Office Staff'),
    ('MGR', 'Managers & Team Leads'),
    ('FIN', 'Finance & Operations'),
    ('HR',  'HR & People Development'),
    ('CS',  'Sales & Customer Service'),
    ('IT',  'IT & Systems Administrators'),
]
for code, name in roles:
    add_bullet(doc, f'{code} — {name}')

doc.add_paragraph()

# ═══════════════════════════════════════════════
# SECTION 2 — METHODOLOGY
# ═══════════════════════════════════════════════
add_heading(doc, '2.  Methodology', level=1)

add_para(doc,
    'The matrix was constructed using a structured expert judgment approach, a recognised method in '
    'management consulting when empirical usage data is not yet available. The process followed three steps.',
    space_after=10)

# Step 1
add_heading(doc, 'Step 1 — Task Analysis per Role', level=2)
add_para(doc,
    'Each role\'s keyActivities and painPoints were extracted from the Chanel Profile Explorer. '
    'These describe what each role does day-to-day and where their biggest productivity friction lies.',
    space_after=6)

add_para(doc, 'Example — Finance & Operations:', bold=True, size=10, space_after=2)
add_bullet(doc, 'Key activity: "Building Excel models for seasonal forecasting"')
add_bullet(doc, 'Pain point: "Building monthly financial reports in Excel takes significant manual effort and is prone to errors"')
doc.add_paragraph()

# Step 2
add_heading(doc, 'Step 2 — Habit-to-Task Matching', level=2)
add_para(doc,
    'Each of the seven Copilot habits was evaluated against the role\'s tasks and pain points. '
    'The habit that most directly reduces the role\'s stated primary pain was assigned rank 1. '
    'The remaining habits were ranked 2–7 in descending order of relevance.',
    space_after=6)

add_para(doc, 'Matching logic for Finance & Operations:', bold=True, size=10, space_after=2)
add_bullet(doc, 'CH7 Data Analyst → Rank 1   (Excel/data analysis directly addresses their #1 pain)')
add_bullet(doc, 'CH6 Document Insights → Rank 2   (reviewing invoices and reports is a core task)')
add_bullet(doc, 'CH2 Research → Rank 3   (validating financial data requires research)')
add_bullet(doc, 'CH5 Content Creator → Rank 4   (drafting financial summaries is secondary)')
add_bullet(doc, 'CH3 Email → Rank 5   (chasing approvals exists but is not the core pain)')
add_bullet(doc, 'CH4 Meeting → Rank 6   (meetings occur but are not a primary productivity blocker)')
add_bullet(doc, 'CH1 Ask-Me-Anything → Rank 7   (least relevant; FIN works with data, not general Q&A)')
doc.add_paragraph()

# Step 3
add_heading(doc, 'Step 3 — Empirical Validation Where Available', level=2)
add_para(doc,
    'For Knowledge Workers (KW) and Managers (MGR), the rankings were cross-checked against published '
    'Microsoft Research findings (see Section 3). For the remaining four roles, the rankings rely on '
    'the structured task analysis in Step 2 and have not yet been validated against empirical usage data.',
    space_after=10)

add_divider(doc)

# ═══════════════════════════════════════════════
# SECTION 3 — CITATIONS
# ═══════════════════════════════════════════════
add_heading(doc, '3.  Citations & Sources', level=1)

# Primary citation
add_heading(doc, 'Primary Citation', level=2)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
r1 = p.add_run('Dillon, E. W., Jaffe, S., Peng, S., & Cambon, A. (2025). ')
r1.font.name = 'Calibri'; r1.font.size = Pt(11)
r2 = p.add_run('Early Impacts of M365 Copilot.')
r2.font.name = 'Calibri'; r2.font.size = Pt(11); r2.font.italic = True
r3 = p.add_run(' Microsoft Research. arXiv:2504.11443.')
r3.font.name = 'Calibri'; r3.font.size = Pt(11)

add_para(doc, 'Available at: https://arxiv.org/abs/2504.11443', size=10, color=(80,80,80), space_after=6)

add_para(doc,
    'This paper reports a randomised controlled trial (RCT) conducted across 56 firms with over 6,000 '
    'workers over a six-month period. It is the primary empirical source used to validate habit rankings '
    'for KW and MGR roles.',
    size=10, space_after=8)

add_heading(doc, 'What This Paper Actually States', level=3)

findings = [
    ('Email — all workers',
     '"Licensees spent 12 fewer minutes reading emails each week, a 7% decrease. Active users saved '
     'more than half an hour per week (18%)."'),
    ('Email — Individual Contributors (KW)',
     '"ICs were the primary drivers of faster reply times, with Copilot users replying to emails '
     '44 minutes faster on average, a 10% improvement."'),
    ('Email — Managers (MGR)',
     '"Managers achieved this mainly by reading fewer emails per week, while individual contributors '
     'mainly saved more time on each email."'),
    ('Document creation — all workers',
     '"Workers with access to Copilot complete Word documents nearly half a day faster on average '
     '(6% improvement). Workers who actively use Copilot complete documents nearly a full day faster (12%)."'),
    ('Meetings — all workers',
     '"Small and statistically insignificant increases in overall meeting time." Teams Copilot was the '
     'most frequently used component but did not show clear productivity gains in this measure.'),
]

for label, quote in findings:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.left_indent = Inches(0.3)
    rb = p.add_run(f'{label}: ')
    rb.font.name = 'Calibri'; rb.font.size = Pt(10); rb.font.bold = True
    rq = p.add_run(quote)
    rq.font.name = 'Calibri'; rq.font.size = Pt(10); rq.font.italic = True

doc.add_paragraph()

# Important note
add_heading(doc, 'Important Scope Note', level=3)
add_para(doc,
    'The Dillon et al. paper does not segment findings by Finance, HR, Sales & Customer Service, or IT roles. '
    'Rankings for FIN, HR, CS, and IT are based on the structured task analysis described in Step 2 above '
    'and should be treated as informed starting estimates, not empirically validated figures.',
    size=10, space_after=10)

add_divider(doc)

# Secondary sources
add_heading(doc, 'Supporting References', level=2)

refs = [
    ('Microsoft Viva Insights — Copilot Adoption Report',
     'Official Microsoft documentation describing role-based segmentation of Copilot usage by job function. '
     'Provides the framework for measuring adoption by department once deployment data is available.',
     'https://learn.microsoft.com/en-us/viva/insights/advanced/analyst/templates/microsoft-365-copilot-adoption'),
    ('Microsoft Inside Track — Unlocking Copilot at the Role Level',
     'Describes Microsoft\'s internal methodology for identifying priority roles and developing "hero scenarios" '
     'based on role size, enthusiasm, and Copilot applicability to repetitive or communication-intensive tasks.',
     'https://www.microsoft.com/insidetrack/blog/unlocking-the-potential-of-copilot-for-microsoft-365-at-the-role-level/'),
    ('Forrester Total Economic Impact of Microsoft 365 Copilot (2024)',
     'Independent study commissioned by Microsoft. Interviewed 8 early-adopter organisations and surveyed 351 '
     'respondents. Reports composite 3-year ROI of 116% and average 9 hours per month saved per active user.',
     'https://tei.forrester.com/go/microsoft/M365Copilot/'),
]

for title, desc, url in refs:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.font.name = 'Calibri'; r.font.size = Pt(11); r.font.bold = True

    add_para(doc, desc, size=10, space_after=2)
    add_para(doc, f'Available at: {url}', size=10, color=(80,80,80), space_after=10)

add_divider(doc)

# ═══════════════════════════════════════════════
# SECTION 4 — RECOMMENDED NEXT STEPS
# ═══════════════════════════════════════════════
add_heading(doc, '4.  Recommended Next Steps', level=1)

add_para(doc,
    'The matrix is a reasonable starting point for guiding Copilot adoption prioritisation. '
    'To move from expert judgment to empirical validation, the following is recommended:',
    space_after=8)

steps = [
    ('Short term (0–3 months)',
     'Use the matrix as-is to guide training prioritisation and use case communication per role.'),
    ('Medium term (3–6 months post-deployment)',
     'Pull actual Copilot feature usage by role from Microsoft Viva Insights Copilot Dashboard. '
     'Compare observed usage patterns against the matrix rankings and update where they diverge.'),
    ('Long term',
     'Conduct a structured user survey (5–10 questions per role group) asking staff to rate which '
     'habits they find most valuable. Use this to build a Chanel-specific, empirically grounded version '
     'of the matrix.'),
]

for label, text in steps:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    rb = p.add_run(f'{label}: ')
    rb.font.name = 'Calibri'; rb.font.size = Pt(11); rb.font.bold = True
    rq = p.add_run(text)
    rq.font.name = 'Calibri'; rq.font.size = Pt(11)

doc.add_paragraph()

# ── Save ──
output_path = 'Role-Habit Relevance Matrix Methodology.docx'
doc.save(output_path)
print(f'Saved: {output_path}')
